# -*- coding: utf-8 -*-
"""FULL-02 学校格式模板核查（plan/01 §11、plan/02 FR-11）。

两件能力，对应两类 check_type：

1. **样式核查（style）**：样式树与角色一致性——三条确定性判据（误报 0 优先，
   仅判结构性可判项，不猜内容）：
   - 章标题（`第N章 …`）必须用 `Heading 1`（否则=标题降级）；
   - 题注段（整段"图/表+编号+题名"，无句读）必须用 `Caption`；
   - 标题样式的段落**不得以句末标点结尾**（句号收尾=正文误用标题样式）。
   数据面=gen_thesis_docx 植入 S1/S2/S3（thesis_truth 已锁）。

2. **模板核查（format）**：docx 节属性/页边距/纸张/页眉页脚/目录 → 核对
   **学校 profile**（四层 ruleset 最高层，`layer4_school_*.json` 用户导入）。
   规则 `detect.param` 指向 `read_doc_format` 抽取的模型路径，`op` ∈
   {eq（带容差）, gte, lte, present}；每条不符合项挂该 profile 条目的条款/来源。

学校 profile 用法见 `data/rulesets/layer4_school_template.json` 的
`import_instructions`：复制为 `layer4_school_<校名>.json`，改 profile=school-<校名>。

依赖：docx 惰性导入（缺装抛 RuntimeError，不静默跳过）。
"""

import re
from pathlib import Path

from taw import engine

EMU_PER_MM = 36000.0

# 题注段判定（与 thesisparse 同口径：整段"图/表+编号+题名"，且无句读标点）
_DASH = "-\uFF0D\u2014\u2013"
_NUM = r"\d{1,2}\s*[%s]?\s*\d{1,2}" % _DASH
_CAPTION_HEAD = re.compile(r"^(图|表)\s*(%s)\s*\S" % _NUM)
_CAPTION_STOP = re.compile(r"[,，。；;：:]")
_CHAP_TITLE = re.compile(r"^第\s*\d{1,2}\s*章[\s\u3000]")
_SENT_END = "。！？!?"


def _ensure_docx():
    try:
        import docx  # noqa: F401
    except ImportError:
        raise RuntimeError("缺少 python-docx：python -m pip install -e './code[docx]'")


# ---- 文档格式模型 ----


def read_doc_format(path):
    """docx → 格式模型（节属性/页边距/纸张/页眉页脚/目录/封面要素）。

    度量一律转 mm（EMU/36000）；抽取**只读**，不改文档。
    """
    _ensure_docx()
    from docx import Document

    doc = Document(str(path))
    sections = []
    for s in doc.sections:
        sections.append({
            "page_width_mm": _mm(s.page_width),
            "page_height_mm": _mm(s.page_height),
            "margin_left_mm": _mm(s.left_margin),
            "margin_right_mm": _mm(s.right_margin),
            "margin_top_mm": _mm(s.top_margin),
            "margin_bottom_mm": _mm(s.bottom_margin),
            "header_present": bool(_header_text(s.header)),
            "footer_present": bool(_header_text(s.footer)),
            "header_text": _header_text(s.header),
            "footer_text": _header_text(s.footer),
        })
    first = sections[0] if sections else {}
    xml = doc.element.xml
    has_toc = ("目录" in [p.text.strip() for p in doc.paragraphs]
               or "TOC" in xml)
    styles = {p.style.name for p in doc.paragraphs}
    cover_text = " ".join(p.text for p in doc.paragraphs[:8])
    return {
        "sections": sections,
        "page_width_mm": first.get("page_width_mm"),
        "page_height_mm": first.get("page_height_mm"),
        "margin_left_mm": first.get("margin_left_mm"),
        "margin_right_mm": first.get("margin_right_mm"),
        "margin_top_mm": first.get("margin_top_mm"),
        "margin_bottom_mm": first.get("margin_bottom_mm"),
        "header_present": first.get("header_present", False),
        "footer_present": first.get("footer_present", False),
        "has_toc": has_toc,
        "has_title_style": "Title" in styles,
        "has_author": "作者" in cover_text,
        "has_school": "学校" in cover_text,
        "has_supervisor": "指导教师" in cover_text or "导师" in cover_text,
    }


def _mm(emu):
    return round(float(emu) / EMU_PER_MM, 2) if emu is not None else None


def _header_text(part):
    try:
        return " ".join(p.text for p in part.paragraphs).strip()
    except (AttributeError, ValueError):
        return ""


# ---- 样式核查（FULL-02 主判定）----


def check_style(parsed, ruleset, doc_name="thesis"):
    """样式核查：章标题层级 / 题注样式 / 正文误用标题样式（三条确定性判据）。"""
    rules = _rule_map(ruleset)
    r_style = _active(rules, "style")
    findings = []
    if r_style is None:
        return findings
    for p in parsed.get("style_tree", []):
        text = p["text"].strip()
        style = p["style"]
        if not text:
            continue
        # 1. 章标题降级
        if _CHAP_TITLE.match(text) and style != "Heading 1":
            findings.append(engine.make_rule_finding(
                r_style, "fail", _loc(doc_name, p["i"], "heading_level"),
                suggestion="章标题改用 Heading 1",
                confidence=0.95,
                detail="章标题使用 {}，应为 Heading 1".format(style),
            ))
            continue
        # 2. 题注样式
        if _is_caption(text) and style != "Caption":
            findings.append(engine.make_rule_finding(
                r_style, "fail", _loc(doc_name, p["i"], "caption_style"),
                suggestion="题注段改用 Caption 样式",
                confidence=0.9,
                detail="题注段落使用 {} 样式，应为 Caption".format(style),
            ))
            continue
        # 3. 正文误用标题样式（标题不以句末标点结尾）
        if style.startswith("Heading") and text[-1] in _SENT_END:
            findings.append(engine.make_rule_finding(
                r_style, "fail", _loc(doc_name, p["i"], "body_style"),
                suggestion="正文段落改用 Normal 样式",
                confidence=0.9,
                detail="正文段落误用 {} 样式".format(style),
            ))
    return findings


def _is_caption(text):
    m = _CAPTION_HEAD.match(text)
    if not m:
        return False
    return _CAPTION_STOP.search(text) is None


# ---- 模板核查（学校 profile 驱动）----


def check_format(doc_format, ruleset, doc_name="thesis"):
    """模板核查：逐条 profile 规则求值，不符合 → finding（挂条款/来源）。"""
    rules = _rule_map(ruleset)
    findings = []
    for r in ruleset["rules"] if ruleset else []:
        if r.get("check_type") != "format":
            continue
        if not engine.is_active(r):
            continue
        detect = r.get("detect") or {}
        param = detect.get("param")
        op = detect.get("op", "eq")
        want = detect.get("value")
        got = _dig(doc_format, param)
        ok, detail = _eval(got, op, want, detect.get("tolerance"))
        if not ok:
            findings.append(engine.make_rule_finding(
                r, "fail", {"doc": doc_name, "para": None, "span": None,
                            "field": param, "label": param},
                suggestion=r.get("fix", {}).get("template") or "按学校模板调整",
                confidence=0.9, detail=detail,
            ))
    return findings


def _dig(model, path):
    cur = model
    for part in str(path).split("."):
        if isinstance(cur, dict) and part in cur:
            cur = cur[part]
        else:
            return None
    return cur


def _eval(got, op, want, tol):
    """规则求值 → (通过?, detail)。求值异常（模型缺字段）视为不通过并如实说明。"""
    if got is None:
        return False, "文档缺少可判定的 {}（无法核对）".format(want)
    tol = 0.0 if tol is None else float(tol)
    if op == "present":
        ok = bool(got) is bool(want)
        return ok, "{}：要求 {}，实际 {}".format("符合" if ok else "不符合",
                                              "存在" if want else "不存在",
                                              "存在" if got else "不存在")
    if op in ("eq", "gte", "lte") and isinstance(want, (int, float)):
        if op == "eq":
            ok = abs(float(got) - float(want)) <= tol
        elif op == "gte":
            ok = float(got) + tol >= float(want)
        else:
            ok = float(got) - tol <= float(want)
        return ok, "{}：要求 {} {}，实际 {}".format(
            "符合" if ok else "不符合", op, want, got)
    if op == "eq":
        ok = got == want
        return ok, "{}：要求为 {}，实际为 {}".format(
            "符合" if ok else "不符合", want, got)
    return False, "未知判定算子 {}（profile 配置错误）".format(op)


# ---- 公共 ----


def _rule_map(ruleset):
    rules = {}
    if ruleset:
        for r in ruleset["rules"]:
            rules.setdefault(r.get("check_type"), r)
    return rules


def _active(rules, check_type):
    r = rules.get(check_type)
    if r is None or not engine.is_active(r):
        return None
    return r


def _loc(doc_name, para, field):
    return {"doc": doc_name, "para": para, "span": None, "field": field}


def audit_school(parsed, doc_format, ruleset, doc_name="thesis"):
    """FULL-02 主入口 → {findings, disclaimers, summary}（样式 + 模板）。"""
    findings = check_style(parsed, ruleset, doc_name)
    if doc_format is not None:
        findings.extend(check_format(doc_format, ruleset, doc_name))
    return {"findings": findings, "disclaimers": list(engine.DISCLAIMERS),
            "summary": engine.summarize(findings)}
